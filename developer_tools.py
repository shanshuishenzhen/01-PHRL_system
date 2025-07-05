# -*- coding: utf-8 -*-
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
import json
import os
import uuid
import random
import argparse
from datetime import datetime, timedelta
import io
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
import subprocess  # 添加这一行
import time  # 添加这一行
import webbrowser  # 添加这一行
import threading  # 添加这一行
import sqlite3  # 添加数据库支持

# -- 解决模块导入路径问题 --
project_root = os.path.abspath(os.path.dirname(__file__))
if project_root not in sys.path:
    sys.path.insert(0, project_root)
# -- 路径问题解决 --

from developer_tools.question_bank_generator import generate_from_excel

# 导入约定管理器
try:
    from common.conventions_manager import get_conventions_manager, apply_conventions
    conventions_manager = get_conventions_manager()
    CONVENTIONS_AVAILABLE = True
    print("✅ 约定管理器加载成功")
except ImportError as e:
    print(f"⚠️ 约定管理器不可用: {e}")
    conventions_manager = None
    CONVENTIONS_AVAILABLE = False

# 全局常量
USER_DATA_FILE = os.path.join(project_root, 'user_management', 'users.json')
# 修改这一行
SAMPLE_QUESTIONS_FILE = os.path.join(project_root, 'question_bank_web', 'questions_sample.xlsx')
TEMPLATE_FILE_NAME = '样例题组题规则模板.xlsx'

class DeveloperTools:
    """
    一个用于开发和测试的辅助工具模块。
    提供快速生成用户、题库、试卷等模拟数据的功能。
    """
    def __init__(self, root):
        self.root = root
        self.root.title("开发工具 - 数据生成助手")
        self.root.geometry("800x650")

        # 样式配置
        style = ttk.Style()
        style.theme_use('clam')
        self.colors = {'danger': '#e74c3c', 'primary': '#3498db', 'success': '#27ae60'}
        style.configure("TNotebook.Tab", padding=[10, 5], font=("Microsoft YaHei", 10))
        style.configure("TLabel", font=("Microsoft YaHei", 10))
        style.configure("TButton", font=("Microsoft YaHei", 10))

        # 主框架
        main_frame = ttk.Frame(self.root, padding=20)
        main_frame.pack(fill=tk.BOTH, expand=True)

        title_label = ttk.Label(main_frame, text="数据生成助手", font=("Microsoft YaHei", 20, "bold"))
        title_label.pack(pady=(0, 20))

        # 标签页
        notebook = ttk.Notebook(main_frame)
        notebook.pack(fill=tk.BOTH, expand=True)
        user_tab = ttk.Frame(notebook)
        question_tab = ttk.Frame(notebook)
        conventions_tab = ttk.Frame(notebook)
        natural_lang_tab = ttk.Frame(notebook)
        validation_tab = ttk.Frame(notebook)
        danger_zone_tab = ttk.Frame(notebook)
        notebook.add(user_tab, text="用户生成")
        notebook.add(question_tab, text="样例题库生成")
        notebook.add(conventions_tab, text="约定管理")
        notebook.add(natural_lang_tab, text="自然语言约定")
        notebook.add(validation_tab, text="验证复核")
        notebook.add(danger_zone_tab, text="危险区域")

        # 初始化变量
        self.uploaded_template_path = tk.StringVar()

        # 初始化各标签页
        self.create_user_generation_tab(user_tab)
        self.create_question_generation_tab(question_tab)
        self.create_conventions_management_tab(conventions_tab)
        self.create_natural_language_tab(natural_lang_tab)
        self.create_validation_tab(validation_tab)
        self.create_danger_zone_tab(danger_zone_tab)
        
    def create_question_generation_tab(self, parent):
        """创建题库生成标签页的UI(新版)"""
        frame = ttk.Frame(parent, padding=20)
        frame.pack(fill=tk.BOTH, expand=True)

        # --- 1. 创建模板 ---
        template_frame = ttk.LabelFrame(frame, text="步骤 1: 获取模板", padding=15)
        template_frame.pack(fill="x", pady=(0, 20))
        ttk.Label(template_frame, text="您可以下载一个空白的Excel模板文件作为开始。").pack(anchor="w")
        create_btn = tk.Button(template_frame, text="下载空白模板", command=self.create_excel_template, bg=self.colors['primary'], fg='white', relief="flat", padx=10)
        create_btn.pack(side="left", pady=10)

        # --- 2. 上传模板 ---
        upload_frame = ttk.LabelFrame(frame, text="步骤 2: 上传您的模板", padding=15)
        upload_frame.pack(fill="x", pady=(0, 20))
        ttk.Label(upload_frame, text="选择一个根据模板填写好的Excel文件 (.xlsx)").pack(anchor="w")
        
        upload_btn_frame = ttk.Frame(upload_frame)
        upload_btn_frame.pack(fill="x", pady=10)
        upload_btn = tk.Button(upload_btn_frame, text="选择文件...", command=self.upload_template, bg=self.colors['primary'], fg='white', relief="flat", padx=10)
        upload_btn.pack(side="left")
        self.uploaded_path_label = ttk.Label(upload_btn_frame, textvariable=self.uploaded_template_path, foreground="gray")
        self.uploaded_path_label.pack(side="left", padx=10)
        self.uploaded_template_path.set("尚未上传文件")

        # --- 3. 生成题库 ---
        generate_frame = ttk.LabelFrame(frame, text="步骤 3: 生成样例题库", padding=15)
        generate_frame.pack(fill="x", pady=(0, 20))
        ttk.Label(generate_frame, text="此操作将根据您上传的Excel文件，生成一个独立的样例题库。").pack(anchor="w")

        # 按钮容器
        btn_container = ttk.Frame(generate_frame)
        btn_container.pack(fill="x", pady=10)

        generate_btn = tk.Button(btn_container, text="生成样例题库", command=self.run_sample_generation, bg=self.colors['success'], fg='white', font=("Microsoft YaHei", 10, "bold"), relief="flat", padx=10, pady=5)
        generate_btn.pack(side="left", padx=(0, 10))

        # --- 4. 管理题库 ---
        manage_frame = ttk.LabelFrame(frame, text="步骤 4: 管理样例题库", padding=15)
        manage_frame.pack(fill="x", pady=(0, 20))
        ttk.Label(manage_frame, text="管理已生成的样例题库，可以查看、编辑或删除。").pack(anchor="w")

        # 管理按钮容器
        manage_btn_container = ttk.Frame(manage_frame)
        manage_btn_container.pack(fill="x", pady=10)

        view_btn = tk.Button(manage_btn_container, text="查看题库", command=self.open_question_bank_manager, bg=self.colors['primary'], fg='white', relief="flat", padx=10)
        view_btn.pack(side="left", padx=(0, 10))

        delete_btn = tk.Button(manage_btn_container, text="删除样例题库", command=self.delete_sample_banks, bg=self.colors['danger'], fg='white', relief="flat", padx=10)
        delete_btn.pack(side="left")
        

    def create_excel_template(self):
        save_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=TEMPLATE_FILE_NAME,
            title="保存模板文件"
        )
        if not save_path:
            return

        data = {
            '题库名称': ['保卫管理员（三级）理论'] + [''] * 22,
            '1级代码': ['A'] + [''] * 7 + ['B'] + [''] * 5 + ['C'] + [''] * 4 + ['D'] + [''] * 3, '1级比重(%)': [20] + [''] * 7 + [10] + [''] * 5 + [45] + [''] * 4 + [25] + [''] * 3,
            '2级代码': ['A-A'] + [''] * 1 + ['A-B'] + [''] * 5 + ['B-A'] + [''] * 1 + ['B-B'] + [''] * 1 + ['B-C'] + [''] * 1 + ['C-A'] + [''] * 2 + ['C-B'] + [''] * 1 + ['D-A'] + [''] * 1 + ['D-B'] + [''] * 1, '2级比重(%)': [5] + [''] * 1 + [15] + [''] * 5 + [4] + [''] * 1 + [3] + [''] * 1 + [3] + [''] * 1 + [25] + [''] * 2 + [20] + [''] * 1 + [12] + [''] * 1 + [13] + [''] * 1,
            '3级代码': ['A-A-A', 'A-A-B', 'A-B-A', 'A-B-B', 'A-B-C', 'A-B-D', 'A-B-E', 'A-B-F', 'B-A-A', 'B-A-B', 'B-B-A', 'B-B-B', 'B-C-A', 'B-C-B', 'C-A-A', 'C-A-B', 'C-A-C', 'C-B-A', 'C-B-B', 'D-A-A', 'D-A-B', 'D-B-A', 'D-B-B'],
            '3级比重(%)': [2.5, 2.5, 2, 3, 2, 3, 1, 4, 2, 2, 2, 1, 1, 2, 7, 8, 10, 10, 10, 6, 6, 6, 7], '知识点数量': [4, 5, 8, 7, 6, 5, 7, 5, 2, 4, 5, 8, 9, 7, 5, 2, 4, 5, 8, 9, 5, 8, 9],
            'B(单选题)': [10, 20, 15, 20, 10, 20, 15, 20, 20, 15, 20, 10, 20, 15, 20, 15, 15, 20, 10, 20, 15, 20, 20], 'G(多选题)': [10, 20, 15, 20, 10, 20, 15, 20, 20, 15, 20, 10, 20, 15, 20, 15, 15, 20, 10, 20, 15, 20, 20],
            'C(判断题)': [8, 18, 13, 18, 8, 18, 13, 18, 18, 13, 18, 8, 18, 13, 18, 13, 13, 18, 8, 18, 13, 18, 18], 'T(填空题)': [5, 15, 10, 15, 5, 15, 10, 15, 15, 10, 15, 5, 15, 10, 15, 10, 10, 15, 5, 15, 10, 15, 15],
            'D(简答题)': [5, 15, 10, 15, 5, 15, 10, 15, 15, 10, 15, 5, 15, 10, 15, 10, 10, 15, 5, 15, 10, 15, 15], 'U(计算题)': [5, 15, 10, 15, 5, 15, 10, 15, 15, 10, 15, 5, 15, 10, 15, 10, 10, 15, 5, 15, 10, 15, 15],
            'W(论述题)': [5, 15, 10, 15, 5, 15, 10, 15, 15, 10, 15, 5, 15, 10, 15, 10, 10, 15, 5, 15, 10, 15, 15], 'E(案例分析题)': [5, 15, 10, 15, 5, 15, 10, 15, 15, 10, 15, 5, 15, 10, 15, 10, 10, 15, 5, 15, 10, 15, 15],
            'F(组合题)': [5, 15, 10, 15, 5, 15, 10, 15, 15, 10, 15, 5, 15, 10, 15, 10, 10, 15, 5, 15, 10, 15, 15]
        }
        df = pd.DataFrame(data)
        try:
            df.to_excel(save_path, index=False)
            messagebox.showinfo("成功", f"模板文件已保存到:\n{save_path}")
        except Exception as e:
            messagebox.showerror("错误", f"保存模板失败: {e}")

    def upload_template(self):
        filepath = filedialog.askopenfilename(
            title="选择组题规则模板",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        if filepath:
            self.uploaded_template_path.set(filepath)

    def run_sample_generation(self):
        template_path = self.uploaded_template_path.get()
        if not template_path or not os.path.exists(template_path):
            messagebox.showwarning("警告", "请先上传一个有效的Excel模板文件。")
            return

        # 询问生成模式
        append_mode = False
        if os.path.exists(SAMPLE_QUESTIONS_FILE):
            choice = messagebox.askyesnocancel(
                "生成模式选择",
                "检测到已存在样例题库文件。\n\n"
                "选择 '是' = 增量生成（如果题库名称不同则追加，相同则替换）\n"
                "选择 '否' = 覆盖模式（完全替换现有文件和数据库）\n"
                "选择 '取消' = 取消操作"
            )
            if choice is None:  # 用户选择取消
                return
            append_mode = choice  # True表示增量模式，False表示覆盖模式

            # 如果是覆盖模式，先清理数据库中的样例题库数据
            if not append_mode:
                try:
                    self._clear_sample_database()
                    print("覆盖模式: 已清理数据库中的样例题库数据")
                except Exception as e:
                    print(f"清理数据库失败: {e}")
                    # 继续执行，不中断生成过程

        try:
            result_data = generate_from_excel(template_path, SAMPLE_QUESTIONS_FILE, append_mode)

            # 处理返回值（兼容新旧版本）
            if len(result_data) == 3:
                total_generated, bank_name, db_success = result_data
            else:
                total_generated, bank_name = result_data
                db_success = False

            mode_text = "增量生成" if append_mode else "覆盖生成"
            db_status = "[成功] 已同步到题库管理模块" if db_success else "[警告] 仅保存为文件"

            result = messagebox.askquestion("成功",
                f"样例题库{mode_text}完毕！\n\n"
                f"共生成 {total_generated} 道题目\n"
                f"文件保存: {SAMPLE_QUESTIONS_FILE}\n"
                f"数据库状态: {db_status}\n\n"
                f"是否要自动启动题库管理系统查看题库？")
            
            if result == 'yes':
                # 检查题库管理应用是否存在
                flask_app_path = os.path.join(os.path.dirname(__file__), 'question_bank_web', 'app.py')
                if not os.path.exists(flask_app_path):
                    messagebox.showerror("错误", "题库管理应用 'question_bank_web/app.py' 未找到！")
                    return
                    
                # 启动Flask应用并导入样例题库
                def start_flask_and_import():
                    try:
                        # 使用Python直接启动，避免弹出cmd窗口
                        import sys
                        flask_dir = os.path.dirname(flask_app_path)

                        # 在后台启动Flask应用
                        process = subprocess.Popen([
                            sys.executable, 'app.py'
                        ], cwd=flask_dir,
                           stdout=subprocess.DEVNULL,
                           stderr=subprocess.DEVNULL,
                           creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)

                        # 等待服务启动
                        time.sleep(3)
                        # 直接打开导入样例题库的URL
                        webbrowser.open_new("http://127.0.0.1:5000/import-sample")

                    except Exception as e:
                        messagebox.showerror("错误", f"启动题库管理失败: {e}")
                
                threading.Thread(target=start_flask_and_import, daemon=True).start()
                
        except Exception as e:
            messagebox.showerror("生成失败", f"生成过程中发生错误: {e}")
            
    def delete_sample_questions(self):
        if not os.path.exists(SAMPLE_QUESTIONS_FILE):
            messagebox.showinfo("提示", "样例题库文件不存在，无需删除。")
            return
            
        if not messagebox.askyesno("严重警告", f"此操作将永久删除以下文件，且无法恢复！\n\n{SAMPLE_QUESTIONS_FILE}\n\n是否确定要继续？"):
            return
            
        try:
            os.remove(SAMPLE_QUESTIONS_FILE)
            messagebox.showinfo("成功", "样例题库已成功删除。")
        except Exception as e:
            messagebox.showerror("删除失败", f"删除文件时发生错误: {e}")

    def open_question_bank_manager(self):
        """打开题库管理模块"""
        try:
            # 启动题库管理模块
            import subprocess
            import sys
            import os

            # 题库管理模块路径
            question_bank_path = os.path.join(os.path.dirname(__file__), "..", "question_bank_web", "app.py")

            if os.path.exists(question_bank_path):
                if os.name == 'nt':  # Windows
                    subprocess.Popen([sys.executable, question_bank_path], shell=True)
                else:  # Linux/Mac
                    subprocess.Popen([sys.executable, question_bank_path])

                messagebox.showinfo("成功", "题库管理模块已启动！\n请在浏览器中查看题库。")
            else:
                messagebox.showerror("错误", "找不到题库管理模块")

        except Exception as e:
            messagebox.showerror("错误", f"启动题库管理模块失败: {e}")

    def delete_sample_banks(self):
        """删除样例题库"""
        try:
            # 1. 删除样例Excel文件
            if os.path.exists(SAMPLE_QUESTIONS_FILE):
                if messagebox.askyesno("确认删除",
                    f"此操作将删除样例题库文件和数据库中的所有样例题库数据。\n\n"
                    f"文件: {SAMPLE_QUESTIONS_FILE}\n\n"
                    f"确定要继续吗？此操作不可撤销！"):

                    try:
                        os.remove(SAMPLE_QUESTIONS_FILE)
                        print(f"已删除样例题库文件: {SAMPLE_QUESTIONS_FILE}")
                    except Exception as e:
                        print(f"删除样例题库文件失败: {e}")
                else:
                    return

            # 2. 删除数据库中的样例题库数据
            # 导入数据库相关模块
            sys.path.append(os.path.join(os.path.dirname(__file__), 'question_bank_web'))

            try:
                from models import QuestionBank, Question
                from sqlalchemy import create_engine, text
                from sqlalchemy.orm import sessionmaker

                # 连接到Web应用数据库（正确的路径）
                db_path = os.path.join(os.path.dirname(__file__), 'question_bank_web', 'questions.db')
                engine = create_engine(f'sqlite:///{db_path}')
                Session = sessionmaker(bind=engine)
                session = Session()

                # 查找所有包含"样例题库"的题库
                sample_banks = session.query(QuestionBank).filter(QuestionBank.name.like('%样例题库%')).all()

                if sample_banks:
                    # 显示找到的样例题库
                    bank_names = [bank.name for bank in sample_banks]
                    bank_list = "\n".join([f"• {name}" for name in bank_names])
                    print(f"找到样例题库:\n{bank_list}")

                    # 强制删除所有样例题库数据
                    session.execute(text("DELETE FROM questions WHERE question_bank_id IN (SELECT id FROM question_banks WHERE name LIKE '%样例题库%')"))
                    session.execute(text("DELETE FROM question_banks WHERE name LIKE '%样例题库%'"))
                    session.commit()

                    print(f"已从数据库删除 {len(sample_banks)} 个样例题库")
                else:
                    print("数据库中没有找到样例题库")

                session.close()
                messagebox.showinfo("成功", "样例题库文件和数据库数据已全部清理完成！")

            except ImportError as e:
                print(f"无法导入数据库模块: {e}")
                # 如果无法导入数据库模块，尝试直接操作SQLite
                import sqlite3

                db_path = os.path.join(os.path.dirname(__file__), 'question_bank_web', 'questions.db')
                if os.path.exists(db_path):
                    conn = sqlite3.connect(db_path)
                    cursor = conn.cursor()

                    # 查找样例题库
                    cursor.execute("SELECT id, name FROM question_banks WHERE name LIKE '%样例题库%'")
                    sample_banks = cursor.fetchall()

                    if sample_banks:
                        bank_list = "\n".join([f"• {name}" for _, name in sample_banks])
                        print(f"找到样例题库:\n{bank_list}")

                        # 删除题库（级联删除题目）
                        cursor.execute("DELETE FROM questions WHERE question_bank_id IN (SELECT id FROM question_banks WHERE name LIKE '%样例题库%')")
                        cursor.execute("DELETE FROM question_banks WHERE name LIKE '%样例题库%'")

                        conn.commit()
                        print(f"已从数据库删除 {len(sample_banks)} 个样例题库")
                    else:
                        print("数据库中没有找到样例题库")

                    conn.close()
                    messagebox.showinfo("成功", "样例题库文件和数据库数据已全部清理完成！")
                else:
                    print(f"数据库文件不存在: {db_path}")
                    messagebox.showinfo("成功", "样例题库文件已删除，数据库文件不存在。")

        except Exception as e:
            messagebox.showerror("错误", f"删除样例题库失败: {e}")
            print(f"删除样例题库失败: {e}")

    def _clear_sample_database(self):
        """清理数据库中的样例题库数据（内部方法）"""
        try:
            # 导入数据库相关模块
            sys.path.append(os.path.join(os.path.dirname(__file__), 'question_bank_web'))

            try:
                from models import QuestionBank, Question
                from sqlalchemy import create_engine, text
                from sqlalchemy.orm import sessionmaker

                # 连接到Web应用数据库
                db_path = os.path.join(os.path.dirname(__file__), 'question_bank_web', 'questions.db')
                engine = create_engine(f'sqlite:///{db_path}')
                Session = sessionmaker(bind=engine)
                session = Session()

                # 删除所有样例题库数据
                session.execute(text("DELETE FROM questions WHERE question_bank_id IN (SELECT id FROM question_banks WHERE name LIKE '%样例题库%')"))
                session.execute(text("DELETE FROM question_banks WHERE name LIKE '%样例题库%'"))
                session.commit()
                session.close()

            except ImportError:
                # 如果无法导入数据库模块，尝试直接操作SQLite
                import sqlite3

                db_path = os.path.join(os.path.dirname(__file__), 'question_bank_web', 'questions.db')
                if os.path.exists(db_path):
                    conn = sqlite3.connect(db_path)
                    cursor = conn.cursor()

                    # 删除样例题库数据
                    cursor.execute("DELETE FROM questions WHERE question_bank_id IN (SELECT id FROM question_banks WHERE name LIKE '%样例题库%')")
                    cursor.execute("DELETE FROM question_banks WHERE name LIKE '%样例题库%'")

                    conn.commit()
                    conn.close()

        except Exception as e:
            print(f"清理数据库失败: {e}")
            raise

    def create_user_generation_tab(self, parent):
        ttk.Label(parent, text="设置要生成的用户数量：", font=("Microsoft YaHei", 12)).pack(anchor="w")
        entry_frame = ttk.Frame(parent)
        entry_frame.pack(fill="x", pady=10)
        self.entries = {}
        user_types = {"考生": "student", "考评员": "evaluator", "管理员": "admin"}
        defaults = {"考生": 88, "考评员": 10, "管理员": 2}
        for i, (label_text, key) in enumerate(user_types.items()):
            ttk.Label(entry_frame, text=f"{label_text}:", width=10).grid(row=i, column=0, padx=5, pady=5)
            self.entries[key] = ttk.Entry(entry_frame)
            self.entries[key].grid(row=i, column=1, padx=5, pady=5)
            self.entries[key].insert(0, str(defaults.get(label_text, 0)))
        generate_btn = tk.Button(parent, text="生成并追加用户", command=self.generate_users_gui, relief="flat", bg=self.colors['primary'], fg='white', font=("Microsoft YaHei", 10, "bold"))
        generate_btn.pack(pady=20)
        info_label = ttk.Label(parent, text="说明：此操作会保留现有用户，并在其基础上追加新用户。\n用户名将以'student_xxx', 'evaluator_xxx' 等形式生成。", justify="left")
        info_label.pack(anchor="w", pady=10)

    def create_danger_zone_tab(self, parent):
        ttk.Label(parent, text="警告：以下操作会永久删除数据，请谨慎使用！", font=("Microsoft YaHei", 12, "bold"), foreground=self.colors['danger']).pack(anchor="w", pady=10)
        btn_frame = ttk.Frame(parent)
        btn_frame.pack(pady=20)
        clear_users_btn = tk.Button(btn_frame, text="清空测试用户", command=self.clear_all_users, bg=self.colors['danger'], fg="white", font=("Microsoft YaHei", 10, "bold"), relief="flat", padx=10, pady=5)
        clear_users_btn.pack(side="left", padx=10)

        delete_questions_btn = tk.Button(btn_frame, text="删除样例题库", command=self.delete_sample_questions, bg=self.colors['danger'], fg="white", font=("Microsoft YaHei", 10, "bold"), relief="flat", padx=10, pady=5)
        delete_questions_btn.pack(side="left", padx=10)

    def generate_users_gui(self):
        try:
            user_counts = {key: int(entry.get()) for key, entry in self.entries.items()}
            total_count = sum(user_counts.values())
            if total_count == 0:
                messagebox.showwarning("提示", "请输入要生成的用户数量。")
                return
            if not messagebox.askyesno("确认操作", f"您确定要生成并追加 {total_count} 个新用户吗？"):
                return
            
            generated_count, _ = _generate_users_logic(**user_counts)
            messagebox.showinfo("成功", f"成功生成并追加了 {generated_count} 个新用户！")

        except ValueError:
            messagebox.showerror("错误", "请输入有效的数字！")
        except Exception as e:
            messagebox.showerror("生成失败", str(e))
    
    def clear_all_users(self):
        if not messagebox.askyesno("严重警告", "此操作将删除所有'考生'、'考评员'和'管理员'角色的用户，且无法恢复！\n\n系统将保留'super_admin'角色的用户。\n是否确定要继续？"):
            return
        try:
            # 1. 更新JSON文件
            if not os.path.exists(USER_DATA_FILE):
                messagebox.showinfo("提示", "用户文件不存在，无需操作。")
                return
            
            with open(USER_DATA_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            initial_count = len(data.get("users", []))
            # 保留 super_admin
            data["users"] = [u for u in data["users"] if u.get("role") == "super_admin"]
            final_count = len(data["users"])
            
            with open(USER_DATA_FILE, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=4)
            
            # 2. 同步更新数据库
            db_path = os.path.join(project_root, 'user_management', 'users.db')
            if os.path.exists(db_path):
                conn = sqlite3.connect(db_path)
                cursor = conn.cursor()
                try:
                    # 删除非超级管理员用户
                    cursor.execute("DELETE FROM users WHERE role != 'super_admin'")
                    conn.commit()
                except Exception as db_error:
                    messagebox.showerror("数据库操作失败", str(db_error))
                finally:
                    conn.close()
                
            messagebox.showinfo("成功", f"操作完成！共删除了 {initial_count - final_count} 个用户。\n数据库已同步更新。")
        except Exception as e:
            messagebox.showerror("操作失败", str(e))

    def create_validation_tab(self, parent):
        """创建验证复核标签页"""
        frame = ttk.Frame(parent, padding=20)
        frame.pack(fill=tk.BOTH, expand=True)

        # 标题
        title_label = ttk.Label(frame, text="题库复核与组卷复核", font=("Microsoft YaHei", 16, "bold"))
        title_label.pack(pady=(0, 20))

        # --- 题库复核区域 ---
        qb_frame = ttk.LabelFrame(frame, text="题库复核（题库生成验证）", padding=15)
        qb_frame.pack(fill="x", pady=(0, 20))

        ttk.Label(qb_frame, text="验证生成的题库是否符合蓝图规则要求", font=("Microsoft YaHei", 10)).pack(anchor="w", pady=(0, 10))

        # 题库复核按钮区域
        qb_btn_frame = ttk.Frame(qb_frame)
        qb_btn_frame.pack(fill="x", pady=5)

        # 自动验证按钮
        auto_validate_btn = tk.Button(qb_btn_frame, text="生成题库并自动验证",
                                    command=self.run_question_bank_generation_with_validation,
                                    bg=self.colors['primary'], fg='white', relief="flat", padx=15, pady=5)
        auto_validate_btn.pack(side="left", padx=(0, 10))

        # 手动验证按钮
        manual_validate_btn = tk.Button(qb_btn_frame, text="手动验证现有题库",
                                      command=self.run_manual_question_bank_validation,
                                      bg=self.colors['success'], fg='white', relief="flat", padx=15, pady=5)
        manual_validate_btn.pack(side="left", padx=(0, 10))

        # 题库验证状态显示
        self.qb_validation_status = tk.StringVar(value="等待验证...")
        status_label = ttk.Label(qb_frame, textvariable=self.qb_validation_status, foreground="gray")
        status_label.pack(anchor="w", pady=(10, 0))

        # 题库验证报告链接区域
        self.qb_report_frame = ttk.Frame(qb_frame)
        self.qb_report_frame.pack(fill="x", pady=(5, 0))

        # --- 组卷复核区域 ---
        paper_frame = ttk.LabelFrame(frame, text="组卷复核（试卷组题验证）", padding=15)
        paper_frame.pack(fill="x", pady=(0, 20))

        ttk.Label(paper_frame, text="分析试卷的三级代码分布和题型统计", font=("Microsoft YaHei", 10)).pack(anchor="w", pady=(0, 10))

        # 组卷复核按钮区域
        paper_btn_frame = ttk.Frame(paper_frame)
        paper_btn_frame.pack(fill="x", pady=5)

        # Web界面按钮
        web_validate_btn = tk.Button(paper_btn_frame, text="打开Web验证界面",
                                   command=self.open_paper_validation_web,
                                   bg=self.colors['primary'], fg='white', relief="flat", padx=15, pady=5)
        web_validate_btn.pack(side="left", padx=(0, 10))

        # 批量验证按钮
        batch_validate_btn = tk.Button(paper_btn_frame, text="批量验证试卷",
                                     command=self.run_batch_paper_validation,
                                     bg=self.colors['success'], fg='white', relief="flat", padx=15, pady=5)
        batch_validate_btn.pack(side="left", padx=(0, 10))

        # 组卷验证状态显示
        self.paper_validation_status = tk.StringVar(value="等待验证...")
        paper_status_label = ttk.Label(paper_frame, textvariable=self.paper_validation_status, foreground="gray")
        paper_status_label.pack(anchor="w", pady=(10, 0))

        # 组卷验证报告链接区域
        self.paper_report_frame = ttk.Frame(paper_frame)
        self.paper_report_frame.pack(fill="x", pady=(5, 0))

        # --- 报告管理区域 ---
        report_frame = ttk.LabelFrame(frame, text="验证报告管理", padding=15)
        report_frame.pack(fill="x", pady=(0, 20))

        ttk.Label(report_frame, text="管理和查看生成的验证报告", font=("Microsoft YaHei", 10)).pack(anchor="w", pady=(0, 10))

        # 报告管理按钮
        report_btn_frame = ttk.Frame(report_frame)
        report_btn_frame.pack(fill="x", pady=5)

        open_reports_btn = tk.Button(report_btn_frame, text="打开报告目录",
                                   command=self.open_reports_directory,
                                   bg=self.colors['primary'], fg='white', relief="flat", padx=15, pady=5)
        open_reports_btn.pack(side="left", padx=(0, 10))

        refresh_reports_btn = tk.Button(report_btn_frame, text="刷新报告列表",
                                      command=self.refresh_validation_reports,
                                      bg=self.colors['success'], fg='white', relief="flat", padx=15, pady=5)
        refresh_reports_btn.pack(side="left", padx=(0, 10))

        # 初始化报告列表
        self.refresh_validation_reports()

    def run_question_bank_generation_with_validation(self):
        """运行题库生成并自动验证"""
        try:
            self.qb_validation_status.set("正在生成题库并验证...")

            # 在新线程中运行，避免阻塞UI
            def run_generation():
                try:
                    # 切换到developer_tools目录
                    original_dir = os.getcwd()
                    developer_tools_dir = os.path.join(project_root, "developer_tools")
                    os.chdir(developer_tools_dir)

                    # 运行题库生成器（会自动触发验证）
                    result = subprocess.run([sys.executable, "question_bank_generator.py"],
                                          capture_output=True, text=True, timeout=300)

                    os.chdir(original_dir)

                    if result.returncode == 0:
                        self.qb_validation_status.set("✅ 题库生成和验证完成")
                        self.refresh_validation_reports()
                    else:
                        self.qb_validation_status.set(f"❌ 生成失败: {result.stderr[:100]}")

                except subprocess.TimeoutExpired:
                    self.qb_validation_status.set("❌ 生成超时")
                except Exception as e:
                    self.qb_validation_status.set(f"❌ 生成错误: {str(e)[:100]}")

            threading.Thread(target=run_generation, daemon=True).start()

        except Exception as e:
            messagebox.showerror("错误", f"启动题库生成失败: {e}")

    def run_manual_question_bank_validation(self):
        """手动验证现有题库"""
        try:
            # 选择蓝图文件
            blueprint_path = filedialog.askopenfilename(
                title="选择蓝图文件",
                filetypes=[("JSON files", "*.json"), ("All files", "*.*")],
                initialdir=os.path.join(project_root, "developer_tools")
            )

            if not blueprint_path:
                return

            # 选择生成的题库文件
            generated_path = filedialog.askopenfilename(
                title="选择生成的题库文件",
                filetypes=[("Excel files", "*.xlsx"), ("JSON files", "*.json"), ("All files", "*.*")],
                initialdir=os.path.join(project_root, "developer_tools")
            )

            if not generated_path:
                return

            self.qb_validation_status.set("正在验证题库...")

            def run_validation():
                try:
                    original_dir = os.getcwd()
                    developer_tools_dir = os.path.join(project_root, "developer_tools")
                    os.chdir(developer_tools_dir)

                    # 运行验证器
                    result = subprocess.run([
                        sys.executable, "question_bank_validator.py",
                        blueprint_path, generated_path
                    ], capture_output=True, text=True, timeout=120)

                    os.chdir(original_dir)

                    if result.returncode == 0:
                        self.qb_validation_status.set("✅ 题库验证完成")
                        self.refresh_validation_reports()
                    else:
                        self.qb_validation_status.set(f"❌ 验证失败: {result.stderr[:100]}")

                except subprocess.TimeoutExpired:
                    self.qb_validation_status.set("❌ 验证超时")
                except Exception as e:
                    self.qb_validation_status.set(f"❌ 验证错误: {str(e)[:100]}")

            threading.Thread(target=run_validation, daemon=True).start()

        except Exception as e:
            messagebox.showerror("错误", f"启动题库验证失败: {e}")

    def open_paper_validation_web(self):
        """打开组卷验证Web界面"""
        try:
            url = "http://localhost:5000/validate-papers"
            webbrowser.open(url)
            self.paper_validation_status.set("✅ 已打开Web验证界面")
        except Exception as e:
            messagebox.showerror("错误", f"打开Web界面失败: {e}")

    def run_batch_paper_validation(self):
        """运行批量试卷验证"""
        try:
            self.paper_validation_status.set("正在批量验证试卷...")

            def run_validation():
                try:
                    original_dir = os.getcwd()
                    question_bank_dir = os.path.join(project_root, "question_bank_web")
                    os.chdir(question_bank_dir)

                    # 运行试卷验证测试
                    result = subprocess.run([sys.executable, "test_paper_validation.py"],
                                          capture_output=True, text=True, timeout=120)

                    os.chdir(original_dir)

                    if result.returncode == 0:
                        self.paper_validation_status.set("✅ 批量验证完成")
                        self.refresh_validation_reports()
                    else:
                        self.paper_validation_status.set(f"❌ 验证失败: {result.stderr[:100]}")

                except subprocess.TimeoutExpired:
                    self.paper_validation_status.set("❌ 验证超时")
                except Exception as e:
                    self.paper_validation_status.set(f"❌ 验证错误: {str(e)[:100]}")

            threading.Thread(target=run_validation, daemon=True).start()

        except Exception as e:
            messagebox.showerror("错误", f"启动批量验证失败: {e}")

    def open_reports_directory(self):
        """打开验证报告目录"""
        try:
            # 打开题库验证报告目录
            qb_reports_dir = os.path.join(project_root, "developer_tools", "validation_reports")
            if os.path.exists(qb_reports_dir):
                os.startfile(qb_reports_dir)

            # 打开试卷验证报告目录
            paper_reports_dir = os.path.join(project_root, "question_bank_web", "paper_validation_reports")
            if os.path.exists(paper_reports_dir):
                os.startfile(paper_reports_dir)

        except Exception as e:
            messagebox.showerror("错误", f"打开报告目录失败: {e}")

    def refresh_validation_reports(self):
        """刷新验证报告列表"""
        try:
            # 清除现有的报告链接
            for widget in self.qb_report_frame.winfo_children():
                widget.destroy()
            for widget in self.paper_report_frame.winfo_children():
                widget.destroy()

            # 题库验证报告
            qb_reports_dir = os.path.join(project_root, "developer_tools", "validation_reports")
            if os.path.exists(qb_reports_dir):
                qb_reports = [f for f in os.listdir(qb_reports_dir) if f.endswith('.xlsx')]
                qb_reports.sort(key=lambda x: os.path.getmtime(os.path.join(qb_reports_dir, x)), reverse=True)

                if qb_reports:
                    ttk.Label(self.qb_report_frame, text="题库验证报告:", font=("Microsoft YaHei", 9, "bold")).pack(anchor="w")
                    for i, report in enumerate(qb_reports[:3]):  # 只显示最新的3个报告
                        report_path = os.path.join(qb_reports_dir, report)
                        link_btn = tk.Button(self.qb_report_frame, text=f"📄 {report}",
                                           command=lambda p=report_path: self.open_report_file(p),
                                           bg="white", fg=self.colors['primary'], relief="flat",
                                           cursor="hand2", anchor="w")
                        link_btn.pack(anchor="w", pady=1)

            # 试卷验证报告
            paper_reports_dirs = [
                os.path.join(project_root, "question_bank_web", "paper_validation_reports"),
                os.path.join(project_root, "question_bank_web", "paper_validation_test_reports")
            ]

            all_paper_reports = []
            for reports_dir in paper_reports_dirs:
                if os.path.exists(reports_dir):
                    reports = [(f, os.path.join(reports_dir, f)) for f in os.listdir(reports_dir) if f.endswith('.xlsx')]
                    all_paper_reports.extend(reports)

            if all_paper_reports:
                all_paper_reports.sort(key=lambda x: os.path.getmtime(x[1]), reverse=True)
                ttk.Label(self.paper_report_frame, text="试卷验证报告:", font=("Microsoft YaHei", 9, "bold")).pack(anchor="w")
                for i, (report_name, report_path) in enumerate(all_paper_reports[:3]):  # 只显示最新的3个报告
                    link_btn = tk.Button(self.paper_report_frame, text=f"📄 {report_name}",
                                       command=lambda p=report_path: self.open_report_file(p),
                                       bg="white", fg=self.colors['primary'], relief="flat",
                                       cursor="hand2", anchor="w")
                    link_btn.pack(anchor="w", pady=1)

        except Exception as e:
            print(f"刷新报告列表失败: {e}")

    def open_report_file(self, file_path):
        """打开验证报告文件"""
        try:
            os.startfile(file_path)
        except Exception as e:
            messagebox.showerror("错误", f"打开报告文件失败: {e}")

    def create_conventions_management_tab(self, parent):
        """创建约定管理标签页"""
        frame = ttk.Frame(parent, padding=20)
        frame.pack(fill=tk.BOTH, expand=True)

        # 标题
        title_label = ttk.Label(frame, text="系统约定管理", font=("Microsoft YaHei", 16, "bold"))
        title_label.pack(pady=(0, 20))

        # 检查约定管理器状态
        if not CONVENTIONS_AVAILABLE:
            error_frame = ttk.Frame(frame)
            error_frame.pack(fill=tk.X, pady=20)

            error_label = ttk.Label(error_frame, text="❌ 约定管理器不可用",
                                  font=("Microsoft YaHei", 12), foreground="red")
            error_label.pack()

            help_label = ttk.Label(error_frame, text="请确保 common/conventions_manager.py 文件存在",
                                 font=("Microsoft YaHei", 10), foreground="gray")
            help_label.pack(pady=(5, 0))
            return

        # 状态显示
        status_frame = ttk.Frame(frame)
        status_frame.pack(fill=tk.X, pady=(0, 10))

        self.conventions_status = tk.StringVar(value="就绪")
        status_label = ttk.Label(status_frame, textvariable=self.conventions_status,
                               font=("Microsoft YaHei", 10), foreground="blue")
        status_label.pack(side=tk.LEFT)

        # 创建主要区域
        main_paned = ttk.PanedWindow(frame, orient=tk.HORIZONTAL)
        main_paned.pack(fill=tk.BOTH, expand=True)

        # 左侧：约定类别列表
        left_frame = ttk.Frame(main_paned)
        main_paned.add(left_frame, weight=1)

        ttk.Label(left_frame, text="约定类别", font=("Microsoft YaHei", 12, "bold")).pack(pady=(0, 10))

        # 约定类别列表框
        listbox_frame = ttk.Frame(left_frame)
        listbox_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        self.conventions_listbox = tk.Listbox(listbox_frame, font=("Microsoft YaHei", 10))
        scrollbar_left = ttk.Scrollbar(listbox_frame, orient=tk.VERTICAL, command=self.conventions_listbox.yview)
        self.conventions_listbox.configure(yscrollcommand=scrollbar_left.set)

        self.conventions_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_left.pack(side=tk.RIGHT, fill=tk.Y)

        self.conventions_listbox.bind('<<ListboxSelect>>', self.on_convention_category_select)

        # 左侧按钮
        left_btn_frame = ttk.Frame(left_frame)
        left_btn_frame.pack(fill=tk.X)

        ttk.Button(left_btn_frame, text="刷新列表", command=self.refresh_conventions_list).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(left_btn_frame, text="新增类别", command=self.add_convention_category).pack(side=tk.LEFT, padx=5)

        # 右侧：约定内容编辑
        right_frame = ttk.Frame(main_paned)
        main_paned.add(right_frame, weight=2)

        ttk.Label(right_frame, text="约定内容编辑", font=("Microsoft YaHei", 12, "bold")).pack(pady=(0, 10))

        # 当前编辑的约定路径
        self.current_convention_path = tk.StringVar()
        path_frame = ttk.Frame(right_frame)
        path_frame.pack(fill=tk.X, pady=(0, 10))
        ttk.Label(path_frame, text="当前路径:").pack(side=tk.LEFT)
        ttk.Label(path_frame, textvariable=self.current_convention_path,
                 font=("Consolas", 10), foreground="blue").pack(side=tk.LEFT, padx=(5, 0))

        # 约定内容文本编辑器
        text_frame = ttk.Frame(right_frame)
        text_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        # 添加滚动条
        self.conventions_text = tk.Text(text_frame, font=("Consolas", 10), wrap=tk.WORD)
        scrollbar_right = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.conventions_text.yview)
        self.conventions_text.configure(yscrollcommand=scrollbar_right.set)

        self.conventions_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_right.pack(side=tk.RIGHT, fill=tk.Y)

        # 右侧按钮
        right_btn_frame = ttk.Frame(right_frame)
        right_btn_frame.pack(fill=tk.X)

        ttk.Button(right_btn_frame, text="保存更改", command=self.save_convention_changes).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(right_btn_frame, text="重置内容", command=self.reset_convention_content).pack(side=tk.LEFT, padx=5)
        ttk.Button(right_btn_frame, text="删除约定", command=self.delete_convention).pack(side=tk.LEFT, padx=5)
        ttk.Button(right_btn_frame, text="验证约定", command=self.validate_conventions).pack(side=tk.LEFT, padx=5)

        # 初始化
        self.refresh_conventions_list()

    def create_natural_language_tab(self, parent):
        """创建自然语言约定标签页"""
        frame = ttk.Frame(parent, padding=20)
        frame.pack(fill=tk.BOTH, expand=True)

        # 标题
        title_label = ttk.Label(frame, text="🗣️ 自然语言约定需求", font=("Microsoft YaHei", 16, "bold"))
        title_label.pack(pady=(0, 10))

        # 说明
        desc_label = ttk.Label(frame,
                              text="用自然语言描述您的约定需求，系统会自动理解并应用到约定模块中",
                              font=("Microsoft YaHei", 10))
        desc_label.pack(pady=(0, 20))

        # 检查约定管理器状态
        if not CONVENTIONS_AVAILABLE:
            error_label = ttk.Label(frame, text="❌ 约定管理器不可用",
                                  font=("Microsoft YaHei", 12), foreground="red")
            error_label.pack(pady=10)
            return

        # 示例提示
        example_frame = ttk.LabelFrame(frame, text="💡 示例", padding=10)
        example_frame.pack(fill=tk.X, pady=(0, 15))

        examples = [
            "• 把超级管理员的密码改成 'admin2024'",
            "• 判断题的选项改成 '对' 和 '错'",
            "• 主题色改成红色，辅助色改成绿色",
            "• 考试时间默认改成90分钟",
            "• 学生默认权限增加查看成绩",
            "• 题库管理端口改成6000"
        ]

        for example in examples:
            ttk.Label(example_frame, text=example, font=("Microsoft YaHei", 9)).pack(anchor="w")

        # 输入区域
        input_label = ttk.Label(frame, text="请用自然语言描述您的约定需求：",
                               font=("Microsoft YaHei", 12, "bold"))
        input_label.pack(anchor="w", pady=(0, 5))

        # 输入文本框
        self.nl_input_text = tk.Text(frame, height=6, font=("Microsoft YaHei", 11))
        self.nl_input_text.pack(fill=tk.X, pady=(0, 10))

        # 按钮区域
        button_frame = ttk.Frame(frame)
        button_frame.pack(fill=tk.X, pady=(0, 10))

        ttk.Button(button_frame, text="🔍 理解需求", command=self.nl_understand_requirement).pack(side=tk.LEFT, padx=(0, 10))
        ttk.Button(button_frame, text="✅ 应用约定", command=self.nl_apply_requirement).pack(side=tk.LEFT, padx=(0, 10))
        ttk.Button(button_frame, text="🗑️ 清空输入", command=self.nl_clear_input).pack(side=tk.LEFT)

        # 理解结果显示
        result_label = ttk.Label(frame, text="理解结果：", font=("Microsoft YaHei", 12, "bold"))
        result_label.pack(anchor="w", pady=(10, 5))

        self.nl_result_text = tk.Text(frame, height=8, font=("Consolas", 10))
        self.nl_result_text.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        # 状态栏
        self.nl_status_var = tk.StringVar(value="就绪")
        status_label = ttk.Label(frame, textvariable=self.nl_status_var,
                               font=("Microsoft YaHei", 9), foreground="blue")
        status_label.pack(side=tk.BOTTOM)

    def refresh_conventions_list(self):
        """刷新约定类别列表"""
        try:
            self.conventions_listbox.delete(0, tk.END)

            if not conventions_manager or not conventions_manager.conventions:
                self.conventions_status.set("❌ 约定配置为空")
                return

            # 获取所有约定类别
            categories = []
            for key in conventions_manager.conventions.keys():
                if key != "system_info":  # 排除系统信息
                    categories.append(key)

            # 添加到列表框
            for category in sorted(categories):
                self.conventions_listbox.insert(tk.END, category)

            self.conventions_status.set(f"✅ 已加载 {len(categories)} 个约定类别")

        except Exception as e:
            self.conventions_status.set(f"❌ 刷新失败: {e}")
            messagebox.showerror("错误", f"刷新约定列表失败: {e}")

    def on_convention_category_select(self, event):
        """约定类别选择事件"""
        try:
            selection = self.conventions_listbox.curselection()
            if not selection:
                return

            category = self.conventions_listbox.get(selection[0])
            self.current_convention_path.set(category)

            # 获取约定内容
            convention_data = conventions_manager.get_convention(category, {})

            # 格式化为可读的文本
            formatted_text = self.format_convention_data(convention_data)

            # 显示在文本编辑器中
            self.conventions_text.delete(1.0, tk.END)
            self.conventions_text.insert(1.0, formatted_text)

            self.conventions_status.set(f"✅ 已加载约定: {category}")

        except Exception as e:
            self.conventions_status.set(f"❌ 加载失败: {e}")
            messagebox.showerror("错误", f"加载约定内容失败: {e}")

    def format_convention_data(self, data, indent=0):
        """格式化约定数据为可读文本"""
        lines = []
        indent_str = "  " * indent

        if isinstance(data, dict):
            for key, value in data.items():
                if isinstance(value, (dict, list)):
                    lines.append(f"{indent_str}{key}:")
                    lines.append(self.format_convention_data(value, indent + 1))
                else:
                    lines.append(f"{indent_str}{key}: {value}")
        elif isinstance(data, list):
            for i, item in enumerate(data):
                if isinstance(item, (dict, list)):
                    lines.append(f"{indent_str}[{i}]:")
                    lines.append(self.format_convention_data(item, indent + 1))
                else:
                    lines.append(f"{indent_str}- {item}")
        else:
            lines.append(f"{indent_str}{data}")

        return "\n".join(lines)

    def parse_convention_text(self, text):
        """解析文本为约定数据结构"""
        try:
            lines = text.strip().split('\n')
            result = {}
            stack = [result]

            for line in lines:
                if not line.strip():
                    continue

                # 计算缩进级别
                indent_level = (len(line) - len(line.lstrip())) // 2
                content = line.strip()

                # 调整栈深度
                while len(stack) > indent_level + 1:
                    stack.pop()

                current_dict = stack[-1]

                if ':' in content and not content.startswith('-'):
                    # 键值对
                    key, value = content.split(':', 1)
                    key = key.strip()
                    value = value.strip()

                    if value:
                        # 尝试转换数据类型
                        if value.lower() in ['true', 'false']:
                            value = value.lower() == 'true'
                        elif value.isdigit():
                            value = int(value)
                        elif value.replace('.', '').isdigit():
                            value = float(value)

                        current_dict[key] = value
                    else:
                        # 空值，准备嵌套结构
                        current_dict[key] = {}
                        stack.append(current_dict[key])
                elif content.startswith('-'):
                    # 列表项
                    item = content[1:].strip()
                    if not isinstance(current_dict, list):
                        # 转换为列表
                        parent_key = list(stack[-2].keys())[-1] if len(stack) > 1 else None
                        if parent_key:
                            stack[-2][parent_key] = []
                            current_dict = stack[-2][parent_key]
                            stack[-1] = current_dict

                    if isinstance(current_dict, list):
                        current_dict.append(item)

            return result

        except Exception as e:
            raise ValueError(f"解析文本失败: {e}")

    def save_convention_changes(self):
        """保存约定更改"""
        try:
            category = self.current_convention_path.get()
            if not category:
                messagebox.showwarning("警告", "请先选择一个约定类别")
                return

            # 获取文本内容
            text_content = self.conventions_text.get(1.0, tk.END).strip()

            if not text_content:
                messagebox.showwarning("警告", "约定内容不能为空")
                return

            # 解析文本为数据结构
            try:
                parsed_data = self.parse_convention_text(text_content)
            except ValueError as e:
                messagebox.showerror("解析错误", f"文本格式错误: {e}")
                return

            # 确认保存
            if not messagebox.askyesno("确认保存", f"确定要保存对约定类别 '{category}' 的更改吗？"):
                return

            # 更新约定
            success = conventions_manager.update_convention(category, parsed_data)

            if success:
                self.conventions_status.set(f"✅ 已保存约定: {category}")
                messagebox.showinfo("成功", f"约定 '{category}' 已成功保存")
                self.refresh_conventions_list()
            else:
                self.conventions_status.set(f"❌ 保存失败: {category}")
                messagebox.showerror("错误", f"保存约定 '{category}' 失败")

        except Exception as e:
            self.conventions_status.set(f"❌ 保存失败: {e}")
            messagebox.showerror("错误", f"保存约定失败: {e}")

    def reset_convention_content(self):
        """重置约定内容"""
        try:
            category = self.current_convention_path.get()
            if not category:
                messagebox.showwarning("警告", "请先选择一个约定类别")
                return

            if not messagebox.askyesno("确认重置", f"确定要重置约定类别 '{category}' 的内容吗？\n这将丢失所有未保存的更改。"):
                return

            # 重新加载原始内容
            convention_data = conventions_manager.get_convention(category, {})
            formatted_text = self.format_convention_data(convention_data)

            self.conventions_text.delete(1.0, tk.END)
            self.conventions_text.insert(1.0, formatted_text)

            self.conventions_status.set(f"✅ 已重置约定: {category}")

        except Exception as e:
            self.conventions_status.set(f"❌ 重置失败: {e}")
            messagebox.showerror("错误", f"重置约定内容失败: {e}")

    def delete_convention(self):
        """删除约定"""
        try:
            category = self.current_convention_path.get()
            if not category:
                messagebox.showwarning("警告", "请先选择一个约定类别")
                return

            if not messagebox.askyesno("确认删除", f"确定要删除约定类别 '{category}' 吗？\n此操作不可恢复！"):
                return

            # 删除约定
            if category in conventions_manager.conventions:
                del conventions_manager.conventions[category]

                # 保存更改
                conventions_manager.save_config()

                # 清空编辑器
                self.conventions_text.delete(1.0, tk.END)
                self.current_convention_path.set("")

                # 刷新列表
                self.refresh_conventions_list()

                self.conventions_status.set(f"✅ 已删除约定: {category}")
                messagebox.showinfo("成功", f"约定类别 '{category}' 已删除")
            else:
                messagebox.showerror("错误", f"约定类别 '{category}' 不存在")

        except Exception as e:
            self.conventions_status.set(f"❌ 删除失败: {e}")
            messagebox.showerror("错误", f"删除约定失败: {e}")

    def add_convention_category(self):
        """添加新的约定类别"""
        try:
            # 弹出输入对话框
            category_name = tk.simpledialog.askstring("新增约定类别", "请输入新约定类别的名称:")

            if not category_name:
                return

            # 验证名称
            if category_name in conventions_manager.conventions:
                messagebox.showerror("错误", f"约定类别 '{category_name}' 已存在")
                return

            if not category_name.replace('_', '').isalnum():
                messagebox.showerror("错误", "约定类别名称只能包含字母、数字和下划线")
                return

            # 创建新约定类别
            new_convention = {
                "description": f"{category_name} 约定配置",
                "created_at": "2025-07-05",
                "example_setting": "example_value"
            }

            # 添加到约定管理器
            success = conventions_manager.update_convention(category_name, new_convention)

            if success:
                self.refresh_conventions_list()

                # 选择新创建的类别
                for i in range(self.conventions_listbox.size()):
                    if self.conventions_listbox.get(i) == category_name:
                        self.conventions_listbox.selection_set(i)
                        self.on_convention_category_select(None)
                        break

                self.conventions_status.set(f"✅ 已创建约定类别: {category_name}")
                messagebox.showinfo("成功", f"约定类别 '{category_name}' 已创建")
            else:
                messagebox.showerror("错误", f"创建约定类别 '{category_name}' 失败")

        except Exception as e:
            self.conventions_status.set(f"❌ 创建失败: {e}")
            messagebox.showerror("错误", f"创建约定类别失败: {e}")

    def validate_conventions(self):
        """验证约定配置"""
        try:
            self.conventions_status.set("🔍 正在验证约定配置...")

            # 验证约定
            errors = conventions_manager.validate_conventions()

            if not errors:
                self.conventions_status.set("✅ 约定配置验证通过")
                messagebox.showinfo("验证成功", "所有约定配置验证通过！")
            else:
                self.conventions_status.set(f"❌ 发现 {len(errors)} 个问题")

                # 显示详细错误信息
                error_text = "发现以下配置问题:\n\n" + "\n".join(f"• {error}" for error in errors)
                messagebox.showerror("验证失败", error_text)

        except Exception as e:
            self.conventions_status.set(f"❌ 验证失败: {e}")
            messagebox.showerror("错误", f"验证约定配置失败: {e}")

    def nl_understand_requirement(self):
        """理解自然语言需求"""
        requirement = self.nl_input_text.get(1.0, tk.END).strip()

        if not requirement:
            messagebox.showwarning("警告", "请输入约定需求")
            return

        self.nl_status_var.set("🔍 正在理解需求...")

        try:
            # 解析自然语言需求
            parsed_result = self.nl_parse_natural_language(requirement)

            if parsed_result:
                # 显示理解结果
                self.nl_result_text.delete(1.0, tk.END)
                self.nl_result_text.insert(1.0, self.nl_format_parsed_result(parsed_result))

                # 保存解析结果
                self.nl_current_parsed_result = parsed_result

                self.nl_status_var.set("✅ 需求理解完成")
                messagebox.showinfo("成功", "需求理解完成！请查看理解结果，确认无误后点击'应用约定'")
            else:
                self.nl_status_var.set("❌ 需求理解失败")
                messagebox.showerror("错误", "无法理解该需求，请尝试用更清晰的语言描述")

        except Exception as e:
            self.nl_status_var.set(f"❌ 理解失败: {e}")
            messagebox.showerror("错误", f"理解需求时出错: {e}")

    def nl_parse_natural_language(self, text):
        """解析自然语言需求"""
        import re

        # 需求解析规则
        parsing_rules = [
            # 超级管理员相关
            {
                "pattern": r"(超级管理员|管理员|admin).*密码.*改成.*['\"]([^'\"]+)['\"]",
                "action": "update_admin_password",
                "category": "authentication"
            },
            {
                "pattern": r"(超级管理员|管理员|admin).*用户名.*改成.*['\"]([^'\"]+)['\"]",
                "action": "update_admin_username",
                "category": "authentication"
            },
            {
                "pattern": r"(超级管理员|管理员|admin).*(隐藏|内置|默认|隐含)",
                "action": "set_admin_hidden",
                "category": "authentication"
            },
            {
                "pattern": r"(切换|设置).*(生产|正式).*模式",
                "action": "switch_to_production",
                "category": "authentication"
            },
            {
                "pattern": r"(软件开发完成|开发完成|完成开发).*用户名.*phrladmin",
                "action": "switch_to_production",
                "category": "authentication"
            },
            {
                "pattern": r"(生成|建议).*(密码|安全密码)",
                "action": "generate_secure_password",
                "category": "authentication"
            },

            # 判断题相关
            {
                "pattern": r"判断题.*选项.*改成.*['\"]([^'\"]+)['\"].*['\"]([^'\"]+)['\"]",
                "action": "update_true_false_options",
                "category": "exam_conventions"
            },

            # UI主题相关
            {
                "pattern": r"主题色.*改成.*(红色|绿色|蓝色|黄色|紫色|橙色|黑色|白色|灰色|#[0-9A-Fa-f]{6})",
                "action": "update_primary_color",
                "category": "ui_conventions"
            },
            {
                "pattern": r"辅助色.*改成.*(红色|绿色|蓝色|黄色|紫色|橙色|黑色|白色|灰色|#[0-9A-Fa-f]{6})",
                "action": "update_secondary_color",
                "category": "ui_conventions"
            },

            # 考试时间相关
            {
                "pattern": r"考试时间.*默认.*改成.*?(\d+).*?分钟",
                "action": "update_exam_duration",
                "category": "exam_conventions"
            },

            # 权限相关
            {
                "pattern": r"(学生|考生).*权限.*增加.*?([^，。\n\r]+)",
                "action": "add_student_permission",
                "category": "authentication"
            },

            # 端口相关
            {
                "pattern": r"(题库管理|question_bank).*端口.*改成.*?(\d+)",
                "action": "update_question_bank_port",
                "category": "network_conventions"
            },
            {
                "pattern": r"(主控台|main_console).*端口.*改成.*?(\d+)",
                "action": "update_main_console_port",
                "category": "network_conventions"
            }
        ]

        # 颜色映射
        color_map = {
            "红色": "#F44336", "绿色": "#4CAF50", "蓝色": "#2196F3",
            "黄色": "#FFEB3B", "紫色": "#9C27B0", "橙色": "#FF9800",
            "黑色": "#000000", "白色": "#FFFFFF", "灰色": "#9E9E9E"
        }

        for rule in parsing_rules:
            match = re.search(rule["pattern"], text, re.IGNORECASE)
            if match:
                result = {
                    "action": rule["action"],
                    "category": rule["category"],
                    "original_text": text,
                    "matched_groups": match.groups()
                }

                # 根据不同的动作处理参数
                if rule["action"] == "update_admin_password":
                    result["new_password"] = match.group(2)
                elif rule["action"] == "update_admin_username":
                    result["new_username"] = match.group(2)
                elif rule["action"] == "set_admin_hidden":
                    result["hidden_type"] = match.group(2)
                elif rule["action"] == "switch_to_production":
                    result["mode"] = "production"
                elif rule["action"] == "generate_secure_password":
                    result["password_type"] = "secure"
                elif rule["action"] == "update_true_false_options":
                    result["option1"] = match.group(1)
                    result["option2"] = match.group(2)
                elif rule["action"] in ["update_primary_color", "update_secondary_color"]:
                    color = match.group(1)
                    result["color"] = color_map.get(color, color)
                elif rule["action"] == "update_exam_duration":
                    result["duration"] = int(match.group(1))
                elif rule["action"] == "add_student_permission":
                    result["permission"] = match.group(2).strip()
                elif rule["action"] in ["update_question_bank_port", "update_main_console_port"]:
                    result["port"] = int(match.group(2))

                return result

        return None

    def nl_format_parsed_result(self, result):
        """格式化解析结果"""
        lines = []
        lines.append("🔍 需求理解结果:")
        lines.append("=" * 40)
        lines.append(f"原始需求: {result['original_text']}")
        lines.append(f"约定类别: {result['category']}")
        lines.append(f"操作类型: {result['action']}")
        lines.append("")

        # 根据不同操作显示具体内容
        if result["action"] == "update_admin_password":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改超级管理员密码为: {result['new_password']}")
            lines.append(f"  • 路径: authentication.super_admin.password")

        elif result["action"] == "update_admin_username":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改超级管理员用户名为: {result['new_username']}")
            lines.append(f"  • 路径: authentication.super_admin.username")

        elif result["action"] == "set_admin_hidden":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 设置超级管理员为{result['hidden_type']}模式")
            lines.append(f"  • 在所有界面中隐藏超级管理员")
            lines.append(f"  • 路径: authentication.super_admin.hidden")

        elif result["action"] == "switch_to_production":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 切换到生产模式")
            lines.append(f"  • 用户名改为: phrladmin")
            lines.append(f"  • 生成系统安全密码")
            lines.append(f"  • 关闭调试模式")

        elif result["action"] == "generate_secure_password":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 生成安全密码")
            lines.append(f"  • 16位长度，包含大小写字母、数字、特殊字符")
            lines.append(f"  • 路径: authentication.super_admin.password")

        elif result["action"] == "update_true_false_options":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改判断题选项为: ['{result['option1']}', '{result['option2']}']")
            lines.append(f"  • 路径: exam_conventions.question_types.true_false.options")

        elif result["action"] == "update_primary_color":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改主题色为: {result['color']}")
            lines.append(f"  • 路径: ui_conventions.theme.primary_color")

        elif result["action"] == "update_secondary_color":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改辅助色为: {result['color']}")
            lines.append(f"  • 路径: ui_conventions.theme.secondary_color")

        elif result["action"] == "update_exam_duration":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改默认考试时间为: {result['duration']} 分钟")
            lines.append(f"  • 路径: exam_conventions.time_limits.default_duration")

        elif result["action"] == "add_student_permission":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 为学生角色添加权限: {result['permission']}")
            lines.append(f"  • 路径: authentication.default_permissions.student")

        elif result["action"] == "update_question_bank_port":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改题库管理端口为: {result['port']}")
            lines.append(f"  • 路径: network_conventions.default_ports.question_bank")

        elif result["action"] == "update_main_console_port":
            lines.append(f"将要执行的操作:")
            lines.append(f"  • 修改主控台端口为: {result['port']}")
            lines.append(f"  • 路径: network_conventions.default_ports.main_console")

        lines.append("")
        lines.append("✅ 确认无误后，请点击'应用约定'按钮")

        return "\n".join(lines)

    def nl_apply_requirement(self):
        """应用约定需求"""
        if not hasattr(self, 'nl_current_parsed_result'):
            messagebox.showwarning("警告", "请先理解需求")
            return

        if not CONVENTIONS_AVAILABLE:
            messagebox.showerror("错误", "约定管理器不可用")
            return

        try:
            result = self.nl_current_parsed_result
            self.nl_status_var.set("🔄 正在应用约定...")

            success = False

            # 根据不同操作应用约定
            if result["action"] == "update_admin_password":
                success = conventions_manager.update_convention(
                    "authentication.super_admin.password",
                    result["new_password"]
                )

            elif result["action"] == "update_admin_username":
                success = conventions_manager.update_convention(
                    "authentication.super_admin.username",
                    result["new_username"]
                )

            elif result["action"] == "set_admin_hidden":
                # 设置超级管理员为隐藏模式
                success = conventions_manager.update_convention(
                    "authentication.super_admin.hidden", True
                ) and conventions_manager.update_convention(
                    "authentication.super_admin.built_in", True
                ) and conventions_manager.update_convention(
                    "authentication.super_admin.implicit", True
                )

            elif result["action"] == "switch_to_production":
                # 切换到生产模式
                try:
                    from common.hidden_super_admin import hidden_super_admin
                    switch_result = hidden_super_admin.switch_to_production_mode()
                    if switch_result["success"]:
                        success = True
                        messagebox.showinfo("生产模式",
                                          f"已切换到生产模式\n新用户名: {switch_result['new_username']}\n新密码: {switch_result['new_password']}\n请妥善保存！")
                    else:
                        success = False
                        messagebox.showerror("错误", switch_result["message"])
                except ImportError:
                    success = False
                    messagebox.showerror("错误", "隐藏超级管理员模块不可用")

            elif result["action"] == "generate_secure_password":
                # 生成安全密码
                try:
                    from common.hidden_super_admin import hidden_super_admin
                    new_password = hidden_super_admin.generate_production_password()
                    if new_password:
                        success = conventions_manager.update_convention(
                            "authentication.super_admin.password", new_password
                        )
                        if success:
                            messagebox.showinfo("密码生成",
                                              f"新密码: {new_password}\n请妥善保存！")
                    else:
                        success = False
                except ImportError:
                    success = False
                    messagebox.showerror("错误", "隐藏超级管理员模块不可用")

            elif result["action"] == "update_true_false_options":
                success = conventions_manager.update_convention(
                    "exam_conventions.question_types.true_false.options",
                    [result["option1"], result["option2"]]
                )

            elif result["action"] == "update_primary_color":
                success = conventions_manager.update_convention(
                    "ui_conventions.theme.primary_color",
                    result["color"]
                )

            elif result["action"] == "update_secondary_color":
                success = conventions_manager.update_convention(
                    "ui_conventions.theme.secondary_color",
                    result["color"]
                )

            elif result["action"] == "update_exam_duration":
                success = conventions_manager.update_convention(
                    "exam_conventions.time_limits.default_duration",
                    result["duration"]
                )

            elif result["action"] == "add_student_permission":
                # 获取当前学生权限
                current_permissions = conventions_manager.get_convention(
                    "authentication.default_permissions.student", []
                )
                if isinstance(current_permissions, list):
                    new_permission = result["permission"]
                    if new_permission not in current_permissions:
                        current_permissions.append(new_permission)
                        success = conventions_manager.update_convention(
                            "authentication.default_permissions.student",
                            current_permissions
                        )
                    else:
                        messagebox.showinfo("提示", f"权限 '{new_permission}' 已存在")
                        success = True

            elif result["action"] == "update_question_bank_port":
                success = conventions_manager.update_convention(
                    "network_conventions.default_ports.question_bank",
                    result["port"]
                )

            elif result["action"] == "update_main_console_port":
                success = conventions_manager.update_convention(
                    "network_conventions.default_ports.main_console",
                    result["port"]
                )

            if success:
                self.nl_status_var.set("✅ 约定应用成功")

                # 清空输入
                self.nl_clear_input()

                # 刷新约定管理标签页（如果存在）
                if hasattr(self, 'refresh_conventions_list'):
                    self.refresh_conventions_list()

                messagebox.showinfo("成功", "约定已成功应用！\n\n系统将在下次启动相关模块时自动使用新约定。")
            else:
                self.nl_status_var.set("❌ 约定应用失败")
                messagebox.showerror("错误", "应用约定失败，请检查约定管理器状态")

        except Exception as e:
            self.nl_status_var.set(f"❌ 应用失败: {e}")
            messagebox.showerror("错误", f"应用约定时出错: {e}")

    def nl_clear_input(self):
        """清空输入"""
        self.nl_input_text.delete(1.0, tk.END)
        self.nl_result_text.delete(1.0, tk.END)
        if hasattr(self, 'nl_current_parsed_result'):
            delattr(self, 'nl_current_parsed_result')
        self.nl_status_var.set("就绪")

def _generate_users_logic(student=0, evaluator=0, admin=0):
    if not os.path.exists(USER_DATA_FILE):
        os.makedirs(os.path.dirname(USER_DATA_FILE), exist_ok=True)
        with open(USER_DATA_FILE, 'w', encoding='utf-8') as f:
            json.dump({"users": []}, f, ensure_ascii=False, indent=4)
    try:
        with open(USER_DATA_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except (json.JSONDecodeError, FileNotFoundError):
        data = {"users": []}

    users = data.get("users", [])
    initial_user_count = len(users)
    generated_students = []
    new_users = []

    role_map = {
        "admin": (admin, "admin", "管理员"),
        "evaluator": (evaluator, "evaluator", "考评员"),
        "student": (student, "student", "考生"),
    }

    for role_key, (count, role_val, role_name) in role_map.items():
        for i in range(count):
            username = f"{role_key}_{random.randint(1000, 99999)}"
            user_entry = _create_user_entry(username, role_val, role_name)
            users.append(user_entry)
            new_users.append(user_entry)
            if role_key == "student":
                generated_students.append(username)
    
    # 1. 更新JSON文件
    data["users"] = users
    with open(USER_DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)
    
    # 2. 同步更新数据库
    db_path = os.path.join(project_root, 'user_management', 'users.db')
    if os.path.exists(db_path):
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        try:
            for user in new_users:
                # 准备用户数据
                user_id = user.get('id')
                id_card = user.get('ID', '')  # 从JSON中的ID字段获取身份证号
                username = user.get('username', '')
                password = user.get('password', user.get('password_hash', '123456'))
                role = user.get('role', 'student')
                status = user.get('status', 'active')
                real_name = user.get('real_name', '')
                email = user.get('email', '')
                phone = user.get('phone', '')
                department = user.get('department', '')
                created_at = user.get('created_at', datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                
                # 插入用户数据
                cursor.execute("""
                    INSERT INTO users (id, id_card, username, password, role, status, real_name, email, phone, department, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (user_id, id_card, username, password, role, status, real_name, email, phone, department, created_at))
            conn.commit()
            print("数据库已同步更新。")
        except Exception as db_error:
            print(f"数据库操作失败: {db_error}")
        finally:
            conn.close()
    
    generated_count = len(users) - initial_user_count
    print(f"成功生成 {admin} 个管理员, {evaluator} 个考评员, {student} 个考生。总计: {generated_count}")
    return generated_count, generated_students

def _create_user_entry(username, role, role_name):
    return {
        "id": str(uuid.uuid4()),
        "username": username,
        "password_hash": "123456", # 设置默认密码为123456方便调试
        "real_name": f"{role_name}_{username}",
        "role": role,
        "department": random.choice(["技术部", "市场部", "人力资源部", "后勤部"]),
        "ID": ''.join([str(random.randint(0, 9)) for _ in range(18)]) # 自动生成18位数字作为身份证号
    }

def prepare_enrollment_files_cli(student_usernames, theory1_count):
    if len(student_usernames) < theory1_count:
        print(f"错误：考生总数 ({len(student_usernames)}) 小于理论一所需人数 ({theory1_count})。")
        return
    random.shuffle(student_usernames)
    theory1_candidates = student_usernames[:theory1_count]
    theory2_candidates = student_usernames[theory1_count:]
    with open("theory1_candidates.txt", 'w', encoding='utf-8') as f:
        for username in theory1_candidates:
            f.write(username + '\n')
    print(f"已生成 'theory1_candidates.txt'，包含 {len(theory1_candidates)} 名考生。")
    with open("theory2_candidates.txt", 'w', encoding='utf-8') as f:
        for username in theory2_candidates:
            f.write(username + '\n')
    print(f"已生成 'theory2_candidates.txt'，包含 {len(theory2_candidates)} 名考生。")

def clear_all_users_cli():
    """命令行方式清空用户"""
    try:
        # 1. 更新JSON文件
        if not os.path.exists(USER_DATA_FILE):
            print("用户文件不存在，无需操作。")
            return
        
        with open(USER_DATA_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        initial_count = len(data.get("users", []))
        # 保留 super_admin
        data["users"] = [u for u in data["users"] if u.get("role") == "super_admin"]
        final_count = len(data["users"])
        
        with open(USER_DATA_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=4)
        
        # 2. 同步更新数据库
        db_path = os.path.join(project_root, 'user_management', 'users.db')
        if os.path.exists(db_path):
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            try:
                # 删除非超级管理员用户
                cursor.execute("DELETE FROM users WHERE role != 'super_admin'")
                conn.commit()
                print("数据库已同步更新。")
            except Exception as db_error:
                print(f"数据库操作失败: {db_error}")
            finally:
                conn.close()
            
        print(f"操作完成！共删除了 {initial_count - final_count} 个用户。")
    except Exception as e:
        print(f"操作失败: {e}")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="PH&RL 开发辅助工具")
    parser.add_argument('task', nargs='?', default='gui', help="要执行的任务: 'gui', 'generate_users', 'clear_all_users'")
    parser.add_argument('--admins', type=int, default=0)
    parser.add_argument('--evaluators', type=int, default=0)
    parser.add_argument('--students', type=int, default=0)
    parser.add_argument('--theory1-count', type=int, default=0)
    args = parser.parse_args()

    if args.task == 'generate_users':
        _, student_list = _generate_users_logic(student=args.students, evaluator=args.evaluators, admin=args.admins)
        if args.theory1_count > 0 and student_list:
            prepare_enrollment_files_cli(student_list, args.theory1_count)
    elif args.task == 'clear_all_users':
        clear_all_users_cli()
    else: # 'gui' or no task
        root = tk.Tk()
        app = DeveloperTools(root)
        root.mainloop()
